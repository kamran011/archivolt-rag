import io

from docx import Document as DocxDocument

from app.core.exceptions import IngestionError
from app.domains.ingestion.parsers.base import BaseParser
from app.domains.ingestion.schemas import DocumentElement, ParsedDocument, ParsedPage


class DocxParser(BaseParser):
    async def parse(self, content: bytes, file_name: str) -> ParsedDocument:
        try:
            doc = DocxDocument(io.BytesIO(content))
        except Exception as e:
            raise IngestionError(f"Failed to open DOCX: {e}")

        elements: list[DocumentElement] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = (para.style.name if para.style else "") or ""
            if style_name.startswith("Heading"):
                level = 1
                digits = "".join(c for c in style_name if c.isdigit())
                if digits:
                    level = min(int(digits), 6)
                elements.append(DocumentElement(type="heading", content=text, level=level))
            elif style_name in ("List Bullet", "List Number", "List Paragraph"):
                elements.append(DocumentElement(type="list", content=text))
            else:
                elements.append(DocumentElement(type="paragraph", content=text))

        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            table_text = "\n".join(rows)
            if table_text.strip():
                elements.append(DocumentElement(type="table", content=table_text))

        title = file_name.rsplit(".", 1)[0]
        for el in elements:
            if el.type == "heading" and el.level == 1:
                title = el.content
                break

        return ParsedDocument(
            title=title,
            pages=[ParsedPage(page_number=1, elements=elements)],
            metadata={"source": "docx"},
        )
