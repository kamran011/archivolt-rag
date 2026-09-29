import io

import pytest
from docx import Document as DocxDocument

from app.domains.ingestion.parsers.docx_parser import DocxParser
from app.domains.ingestion.parsers.markdown_parser import MarkdownParser
from app.domains.ingestion.parsers.text_parser import TextParser


@pytest.mark.asyncio
async def test_markdown_parser_extracts_headings_and_levels():
    content = b"""# Title

Some intro text.

## Section A

Content of section A.

### Subsection A.1

Deep content.

## Section B

Content of section B.
"""
    parser = MarkdownParser()
    parsed = await parser.parse(content, "doc.md")

    headings = [el for el in parsed.all_elements() if el.type == "heading"]
    assert len(headings) == 4
    assert headings[0].content == "Title"
    assert headings[0].level == 1
    assert headings[2].content == "Subsection A.1"
    assert headings[2].level == 3
    assert parsed.title == "Title"


@pytest.mark.asyncio
async def test_markdown_parser_extracts_code_blocks():
    content = b"""# Doc

```python
def foo():
    return 42
```

After code.
"""
    parser = MarkdownParser()
    parsed = await parser.parse(content, "doc.md")

    code_elements = [el for el in parsed.all_elements() if el.type == "code"]
    assert len(code_elements) == 1
    assert "def foo():" in code_elements[0].content


@pytest.mark.asyncio
async def test_markdown_parser_extracts_lists():
    content = b"""# Doc

- item one
- item two
- item three
"""
    parser = MarkdownParser()
    parsed = await parser.parse(content, "doc.md")

    list_elements = [el for el in parsed.all_elements() if el.type == "list"]
    assert len(list_elements) == 1
    assert "item one" in list_elements[0].content


@pytest.mark.asyncio
async def test_text_parser_splits_paragraphs():
    content = b"First paragraph text.\n\nSecond paragraph text.\n\nThird paragraph."
    parser = TextParser()
    parsed = await parser.parse(content, "notes.txt")

    paragraphs = [el for el in parsed.all_elements() if el.type == "paragraph"]
    assert len(paragraphs) == 3


@pytest.mark.asyncio
async def test_text_parser_detects_heading_heuristic():
    content = b"SECTION ONE:\n\nSome body text follows here."
    parser = TextParser()
    parsed = await parser.parse(content, "notes.txt")

    headings = [el for el in parsed.all_elements() if el.type == "heading"]
    assert len(headings) == 1
    assert headings[0].content == "SECTION ONE:"


@pytest.mark.asyncio
async def test_docx_parser_extracts_headings_and_paragraphs():
    doc = DocxDocument()
    doc.add_heading("Document Title", level=1)
    doc.add_paragraph("Intro paragraph.")
    doc.add_heading("Section One", level=2)
    doc.add_paragraph("Section one content.")

    buffer = io.BytesIO()
    doc.save(buffer)
    content = buffer.getvalue()

    parser = DocxParser()
    parsed = await parser.parse(content, "test.docx")

    headings = [el for el in parsed.all_elements() if el.type == "heading"]
    assert len(headings) == 2
    assert headings[0].content == "Document Title"
    assert headings[0].level == 1
    assert headings[1].level == 2
    assert parsed.title == "Document Title"


@pytest.mark.asyncio
async def test_docx_parser_extracts_tables():
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header1"
    table.cell(0, 1).text = "Header2"
    table.cell(1, 0).text = "Val1"
    table.cell(1, 1).text = "Val2"

    buffer = io.BytesIO()
    doc.save(buffer)
    content = buffer.getvalue()

    parser = DocxParser()
    parsed = await parser.parse(content, "test.docx")

    tables = [el for el in parsed.all_elements() if el.type == "table"]
    assert len(tables) == 1
    assert "Header1" in tables[0].content
    assert "Val2" in tables[0].content
